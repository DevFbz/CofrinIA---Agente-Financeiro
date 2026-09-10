from __future__ import annotations

import csv
import io
from collections.abc import Iterable
from datetime import date, datetime
from typing import Any

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill

_CATEGORY_LABELS = {
    "alimentacao": "Alimentação",
    "transporte": "Transporte",
    "moradia": "Moradia",
    "lazer": "Lazer",
    "saude": "Saúde",
    "salario": "Salário",
    "outros": "Outros",
}


def _date_label(value: date | datetime | str) -> str:
    if isinstance(value, datetime):
        value = value.date()
    if isinstance(value, date):
        return value.strftime("%d/%m/%Y")
    return str(value)


def _money_label(amount_cents: int) -> str:
    value = amount_cents / 100
    return f"{value:,.2f}".replace(",", "X").replace(".", ",").replace("X", ".")


def build_transactions_csv(rows: Iterable[dict[str, Any]]) -> bytes:
    buffer = io.StringIO(newline="")
    writer = csv.writer(buffer, delimiter=";", lineterminator="\n")
    writer.writerow(("Data", "Tipo", "Valor", "Descrição", "Categoria", "Pagamento"))
    for row in rows:
        transaction_type = "Receita" if row.get("type") == "income" else "Despesa"
        writer.writerow(
            (
                _date_label(row["occurred_on"]),
                transaction_type,
                _money_label(int(row["amount_cents"])),
                row.get("description") or "não informado",
                _CATEGORY_LABELS.get(row.get("category"), "Outros"),
                row.get("payment_method") or "não informado",
            )
        )
    return buffer.getvalue().encode("utf-8-sig")


def build_transactions_xlsx(rows: Iterable[dict[str, Any]]) -> bytes:
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Lançamentos"
    headers = ("Data", "Tipo", "Valor", "Descrição", "Categoria", "Pagamento")
    sheet.append(headers)
    for cell in sheet[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="166534")
    for row in rows:
        sheet.append(
            (
                row["occurred_on"],
                "Receita" if row.get("type") == "income" else "Despesa",
                int(row["amount_cents"]) / 100,
                row.get("description") or "não informado",
                _CATEGORY_LABELS.get(row.get("category"), "Outros"),
                row.get("payment_method") or "não informado",
            )
        )
    for cell in sheet["C"][1:]:
        cell.number_format = 'R$ #,##0.00'
    sheet.freeze_panes = "A2"
    sheet.auto_filter.ref = sheet.dimensions
    for column, width in {"A": 14, "B": 12, "C": 16, "D": 32, "E": 18, "F": 22}.items():
        sheet.column_dimensions[column].width = width
    output = io.BytesIO()
    workbook.save(output)
    return output.getvalue()
